#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成「2026 年度兼职业务盘点」汇报 PPT（管理层版）。

数据来源：flutter-suixinpai 与 yatai_hub 两个仓库的 git 记录，均已核验。
运行：/tmp/pptx-venv/bin/python docs/build_parttime_ppt.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── 配色（dataviz 参考调色板，已过 validate_palette.js）─────────────
BLUE     = "2a78d6"   # series-1 主色
ORANGE   = "eb6834"   # series-2
AQUA     = "1baf7a"   # series-3
GOOD     = "0ca30c"   # status good
WARNING  = "fab219"   # status warning
CRITICAL = "d03b3b"   # status critical
INK      = "0b0b0b"   # primary ink
INK2     = "52514e"   # secondary ink
MUTED    = "898781"   # muted
LINE     = "e1e0d9"   # hairline
SURFACE  = "fcfcfb"   # chart surface
WHITE    = "ffffff"

FONT = "PingFang SC"
W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.75)


def set_font(run, size=14, bold=False, color=INK, name=FONT):
    """设置字体，含中文 eastAsian typeface（python-pptx 默认只设 latin）。"""
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def textbox(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=14, bold=False, color=INK, space_before=0,
         space_after=6, first=False, align=PP_ALIGN.LEFT, line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    set_font(r, size=size, bold=bold, color=color)
    return p


def rect(slide, x, y, w, h, fill=None, line_color=None, line_w=1.0,
         shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line_color:
        sh.line.color.rgb = RGBColor.from_string(line_color)
        sh.line.width = Pt(line_w)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def page_head(slide, title, kicker=None, idx=None):
    """统一页头：小眉线 + 标题 + 右下页码。"""
    y = Inches(0.52)
    if kicker:
        tf = textbox(slide, MARGIN, y, Inches(11.8), Inches(0.28))
        para(tf, kicker, size=11.5, bold=True, color=BLUE, first=True, space_after=0)
        y = Inches(0.85)
    tf = textbox(slide, MARGIN, y, Inches(11.8), Inches(0.55))
    para(tf, title, size=27, bold=True, color=INK, first=True, space_after=0)
    rect(slide, MARGIN, Inches(1.52), Inches(11.83), Emu(9525), fill=LINE)
    if idx:
        tf = textbox(slide, Inches(12.1), Inches(6.92), Inches(0.55), Inches(0.3),
                     align=PP_ALIGN.RIGHT)
        para(tf, str(idx), size=10, color=MUTED, first=True, space_after=0)


def footnote(slide, text):
    tf = textbox(slide, MARGIN, Inches(6.9), Inches(10.9), Inches(0.34))
    para(tf, text, size=9.5, color=MUTED, first=True, space_after=0)


prs = Presentation()
prs.slide_width, prs.slide_height = W, H

# ═══ P1 封面 ════════════════════════════════════════════════════════
s = blank(prs)
rect(s, 0, 0, W, H, fill=SURFACE)
rect(s, 0, 0, Inches(0.14), H, fill=BLUE)
tf = textbox(s, Inches(1.1), Inches(2.15), Inches(11), Inches(0.4))
para(tf, "2026 年度项目盘点 · 专题汇报", size=13, bold=True, color=BLUE, first=True)
tf = textbox(s, Inches(1.1), Inches(2.62), Inches(11), Inches(1.0))
para(tf, "兼职业务盘点与下半年计划", size=42, bold=True, color=INK, first=True)
rect(s, Inches(1.1), Inches(3.92), Inches(1.5), Inches(0.035), fill=BLUE)
tf = textbox(s, Inches(1.1), Inches(4.25), Inches(11), Inches(0.9))
para(tf, "覆盖范围：SZST 新质 OA 兼职业务体系（P08） + 亚太HUB 兼职专用端（P15）",
     size=15, color=INK2, first=True, space_after=8)
para(tf, "盘点周期：2026-01-01 ~ 2026-08-03", size=15, color=INK2, space_after=0)
tf = textbox(s, Inches(1.1), Inches(6.35), Inches(11), Inches(0.7))
para(tf, "数据来源：flutter-suixinpai 与 yatai_hub 两个代码仓库的提交记录、版本记录、源码现状",
     size=10.5, color=MUTED, first=True, space_after=4)
para(tf, "全部数据可客观核验；凡属反推口径均已在末页标注    |    编制日期 2026-08-03",
     size=10.5, color=MUTED, space_after=0)

# ═══ P2 一页看全 ════════════════════════════════════════════════════
s = blank(prs)
page_head(s, "一页看全：兼职业务今年做成了什么", kicker="核心结论", idx=2)

tf = textbox(s, MARGIN, Inches(1.78), Inches(11.83), Inches(0.5))
para(tf, "兼职业务是本年度从 0 到 1 最快形成闭环的新业务：5 月立项，7 月底完成全链路含二期。"
         "但承载它的独立 App 尚未正式发版。",
     size=16, bold=True, color=INK, first=True, space_after=0, line=1.35)

cards = [
    ("2 个",   "兼职相关立项项目",   "业务体系 + 独立端",       BLUE),
    ("153 个", "投入提交数",         "P08 93 + P15 60",         BLUE),
    ("6 / 6",  "业务闭环环节交付",   "准入→接单→打卡→录资料→结算→报销", GOOD),
    ("0 次",   "独立端正式发版",     "仍停在 1.0.0+1，0 个 tag", CRITICAL),
]
cw, gap = Inches(2.79), Inches(0.235)
for i, (num, label, sub, color) in enumerate(cards):
    x = MARGIN + i * (cw + gap)
    rect(s, x, Inches(2.62), cw, Inches(1.62), fill=WHITE, line_color=LINE)
    rect(s, x, Inches(2.62), cw, Inches(0.055), fill=color)
    tf = textbox(s, x + Inches(0.22), Inches(2.85), cw - Inches(0.44), Inches(0.6))
    para(tf, num, size=31, bold=True, color=color, first=True, space_after=0)
    tf = textbox(s, x + Inches(0.22), Inches(3.48), cw - Inches(0.44), Inches(0.28))
    para(tf, label, size=12.5, bold=True, color=INK, first=True, space_after=3)
    para(tf, sub, size=9.5, color=MUTED, space_after=0)

tf = textbox(s, MARGIN, Inches(4.55), Inches(11.83), Inches(0.32))
para(tf, "三句话概括", size=12.5, bold=True, color=BLUE, first=True, space_after=0)
rows = [
    ("做成了", "兼职全链路闭环上线并随母 App 发版 5 次；报销单在一期内吃掉原定二期范围", GOOD),
    ("在推进", "亚太HUB 兼职专用端已从 OA 剥离完成，功能与品牌就绪，等发版", WARNING),
    ("要注意", "独立端未发版 + 素材上传链路零验证 + 无业务埋点，成效说不出数字", CRITICAL),
]
for i, (tag, text, color) in enumerate(rows):
    y = Inches(4.95) + i * Inches(0.6)
    rect(s, MARGIN, y, Inches(0.05), Inches(0.42), fill=color)
    tf = textbox(s, MARGIN + Inches(0.22), y + Inches(0.02), Inches(1.0), Inches(0.3))
    para(tf, tag, size=12.5, bold=True, color=color, first=True, space_after=0)
    tf = textbox(s, MARGIN + Inches(1.25), y + Inches(0.02), Inches(10.4), Inches(0.38))
    para(tf, text, size=13, color=INK2, first=True, space_after=0, line=1.2)
footnote(s, "状态色均配文字标签，不以颜色单独承载含义。")

# ═══ P3 立项概况 ════════════════════════════════════════════════════
s = blank(prs)
page_head(s, "兼职相关立项项目共 2 个", kicker="一、立项概况", idx=3)

projs = [
    (BLUE, "P08", "兼职业务体系", "2026-05 立项 · 93 个提交 · 母仓库",
     ["预设交付口径（反推）：准入审核 → 接单 → 打卡 → 录入客户资料 →",
      "钱包结算 → 报销，全链路闭环",
      "承载于 SZST 新质 OA，与客服 / 商务 / 管理员同一 App 内协同",
      "状态：已交付，含二期"],
     "已交付", GOOD),
    (ORANGE, "P15", "亚太HUB 兼职专用端", "2026-07-13 建仓 · 60 个提交 · 独立仓库",
     ["预设交付口径（反推）：从 OA 剥离「兼职工作流 + 素材上传 +",
      "任务群聊」三条链路，独立品牌 / 包名 / 推送 / 合规主体",
      "目标：形成兼职人员专用、可独立发版的 App",
      "状态：功能与品牌就绪，未正式发版"],
     "在建", WARNING),
]
for i, (color, code, name, meta, bullets, badge, bcolor) in enumerate(projs):
    x = MARGIN + i * Inches(6.05)
    rect(s, x, Inches(1.85), Inches(5.78), Inches(3.32), fill=WHITE, line_color=LINE)
    rect(s, x, Inches(1.85), Inches(0.055), Inches(3.32), fill=color)
    tf = textbox(s, x + Inches(0.3), Inches(2.08), Inches(1.0), Inches(0.3))
    para(tf, code, size=17, bold=True, color=color, first=True, space_after=0)
    tf = textbox(s, x + Inches(1.15), Inches(2.1), Inches(3.2), Inches(0.32))
    para(tf, name, size=17, bold=True, color=INK, first=True, space_after=0)
    bx = x + Inches(4.42)
    rect(s, bx, Inches(2.12), Inches(1.05), Inches(0.32), fill=bcolor)
    tf = textbox(s, bx, Inches(2.18), Inches(1.05), Inches(0.26), align=PP_ALIGN.CENTER)
    para(tf, badge, size=11, bold=True, color=WHITE, first=True, space_after=0)
    tf = textbox(s, x + Inches(0.3), Inches(2.56), Inches(5.2), Inches(0.26))
    para(tf, meta, size=10.5, color=MUTED, first=True, space_after=0)
    tf = textbox(s, x + Inches(0.3), Inches(2.98), Inches(5.2), Inches(1.9))
    for j, b in enumerate(bullets):
        para(tf, b, size=12, color=INK2, first=(j == 0), space_after=5, line=1.15)

tf = textbox(s, MARGIN, Inches(5.42), Inches(11.83), Inches(0.32))
para(tf, "另有一项未正式立项、却持续消耗产能的工作", size=12.5, bold=True,
     color=CRITICAL, first=True, space_after=6)
rect(s, MARGIN, Inches(5.82), Inches(11.83), Inches(0.88), fill=WHITE, line_color=LINE)
tf = textbox(s, MARGIN + Inches(0.28), Inches(6.0), Inches(11.3), Inches(0.6))
para(tf, "双端人工同步：亚太HUB 60 个提交中 17 个（28%）标题明确写「同步 flutter-suixinpai」",
     size=13, bold=True, color=INK, first=True, space_after=4)
para(tf, "两个 App 的群聊、钱包结算、注册页需保持同一份代码，目前全靠人工双向搬运，"
         "无脚本、无对齐清单、无差异校验。",
     size=11.5, color=INK2, space_after=0)
footnote(s, "「预设交付口径」为依据首个提交范围与发版节点反推，非立项时签署的量化 KPI —— 仓库不承载项目管理数据。")

# ═══ P4 业务闭环全景 ════════════════════════════════════════════════
s = blank(prs)
page_head(s, "兼职业务闭环：6 个环节全部交付", kicker="一、立项概况", idx=4)

steps = [
    ("准入", "注册审核\n护照 / 收款码"),
    ("接单", "接单大厅\n任务详情"),
    ("打卡", "上下班双接口\n计薪工时"),
    ("录资料", "客户 / 订单\n审核驳回重交"),
    ("结算", "钱包 / 流水\n提现 / 月结"),
    ("报销", "独立报销单\n编辑 / 取消"),
]
bw, bgap = Inches(1.72), Inches(0.265)
total = len(steps) * bw + (len(steps) - 1) * bgap
sx = (W - total) // 2
for i, (name, desc) in enumerate(steps):
    x = sx + i * (bw + bgap)
    rect(s, x, Inches(2.05), bw, Inches(1.62), fill=WHITE, line_color=BLUE, line_w=1.3)
    tf = textbox(s, x, Inches(2.28), bw, Inches(0.34), align=PP_ALIGN.CENTER)
    para(tf, name, size=17, bold=True, color=BLUE, first=True, space_after=0)
    tf = textbox(s, x, Inches(2.76), bw, Inches(0.75), align=PP_ALIGN.CENTER)
    for j, ln in enumerate(desc.split("\n")):
        para(tf, ln, size=11, color=INK2, first=(j == 0), space_after=2, align=PP_ALIGN.CENTER)
    tf = textbox(s, x, Inches(3.78), bw, Inches(0.26), align=PP_ALIGN.CENTER)
    para(tf, "已交付", size=10.5, bold=True, color=GOOD, first=True, space_after=0,
         align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        ax = x + bw + Inches(0.055)
        tf = textbox(s, ax, Inches(2.66), Inches(0.16), Inches(0.3), align=PP_ALIGN.CENTER)
        para(tf, "▸", size=15, bold=True, color=MUTED, first=True, space_after=0,
             align=PP_ALIGN.CENTER)

tf = textbox(s, MARGIN, Inches(4.42), Inches(11.83), Inches(0.32))
para(tf, "闭环之外，另交付两项支撑能力", size=12.5, bold=True, color=BLUE, first=True,
     space_after=8)
extra = [
    ("跨角色群聊", "兼职与责任客服、管理员在任务群 / 科室群 / 私聊三态内沟通；"
                   "群内可直接查医院信息与客户资料，沟通与查数同场景完成"),
    ("推送到达保障", "按 App 分投、RID 迟到补报、高优先级消息渠道；"
                     "覆盖首装弱网与 iOS 首启授权晚给等边缘场景"),
]
for i, (name, desc) in enumerate(extra):
    y = Inches(4.82) + i * Inches(0.82)
    rect(s, MARGIN, y, Inches(11.83), Inches(0.68), fill=WHITE, line_color=LINE)
    rect(s, MARGIN, y, Inches(0.05), Inches(0.68), fill=AQUA)
    tf = textbox(s, MARGIN + Inches(0.26), y + Inches(0.11), Inches(2.0), Inches(0.3))
    para(tf, name, size=13, bold=True, color=INK, first=True, space_after=0)
    tf = textbox(s, MARGIN + Inches(2.3), y + Inches(0.09), Inches(9.3), Inches(0.5))
    para(tf, desc, size=11.5, color=INK2, first=True, space_after=0, line=1.15)

# ═══ P5 实施进度：五阶段 ════════════════════════════════════════════
s = blank(prs)
page_head(s, "实施进度：五阶段递进，2.5 个月完成闭环", kicker="二、项目实施进度", idx=5)

phases = [
    ("5 月下旬", "一阶段 · 准入", "待审核拦截页、资料上传页、注册随表单提交护照与收款码、简易合同"),
    ("5 月底",   "二阶段 · 接单", "接单大厅、任务详情与状态流转、「我已接单」接入 my_tasks、到院信息采集"),
    ("6 月",     "三阶段 · 打卡与资料", "日打卡上线；任务数据录入；数据审核与驳回重编辑；展示午休扣减后的计薪工时"),
    ("6 月中",   "四阶段 · 钱包结算", "钱包 / 结算交付；收款账号管理与结算津贴；结算单详情展示折合人民币与汇率"),
    ("7 月下旬", "五阶段 · 报销单", "打卡改 clock_in / clock_out 双接口；独立报销单全链路（提交 / 展示 / 编辑 / 取消）"),
]
rect(s, MARGIN + Inches(1.5), Inches(1.95), Emu(19050), Inches(4.2), fill=LINE)
for i, (when, name, desc) in enumerate(phases):
    y = Inches(1.95) + i * Inches(0.86)
    tf = textbox(s, MARGIN, y + Inches(0.06), Inches(1.3), Inches(0.28),
                 align=PP_ALIGN.RIGHT)
    para(tf, when, size=12, bold=True, color=MUTED, first=True, space_after=0,
         align=PP_ALIGN.RIGHT)
    d = Inches(0.17)
    rect(s, MARGIN + Inches(1.5) - d // 2 + Emu(9525), y + Inches(0.1), d, d,
         fill=BLUE, shape=MSO_SHAPE.OVAL)
    tf = textbox(s, MARGIN + Inches(1.85), y, Inches(3.0), Inches(0.3))
    para(tf, name, size=14, bold=True, color=INK, first=True, space_after=0)
    tf = textbox(s, MARGIN + Inches(1.85), y + Inches(0.32), Inches(9.7), Inches(0.42))
    para(tf, desc, size=11.5, color=INK2, first=True, space_after=0, line=1.15)

rect(s, MARGIN, Inches(6.32), Inches(11.83), Inches(0.5), fill=WHITE, line_color=GOOD)
tf = textbox(s, MARGIN + Inches(0.26), Inches(6.43), Inches(11.3), Inches(0.3))
para(tf, "亮点：报销单在一期内吃掉原定二期范围 —— 7-27 当日拿到 /v1/parttime_expense 契约，"
         "当日即落地编辑 / 重提 / 取消，说明与后端协作链路通畅。",
     size=11.5, bold=True, color=INK, first=True, space_after=0)

# ═══ P6 发版里程碑 ══════════════════════════════════════════════════
s = blank(prs)
page_head(s, "兼职功能已随母 App 发版 5 次", kicker="二、项目实施进度", idx=6)

rels = [
    ("1.2.5+25", "05-27", "兼职业务体系 + 推送体系上线", GOOD),
    ("1.3.0+30", "06-11", "兼职打卡与数据录入完善", GOOD),
    ("1.3.2+32", "06-15", "兼职钱包结算上线", GOOD),
    ("1.3.4+34", "06-17", "群聊 IM 上线（兼职与客服跨角色沟通）", GOOD),
    ("1.4.4+44", "07-30", "独立报销单 + 私聊通讯录（当前生产版本）", GOOD),
]
tf = textbox(s, MARGIN, Inches(1.78), Inches(11.83), Inches(0.3))
para(tf, "SZST 新质 OA（母 App）· 全年约 42 次发版，平均约 5 天一版",
     size=12.5, bold=True, color=BLUE, first=True, space_after=0)
for i, (ver, date, note, color) in enumerate(rels):
    y = Inches(2.2) + i * Inches(0.6)
    rect(s, MARGIN, y, Inches(11.83), Inches(0.5), fill=WHITE, line_color=LINE)
    rect(s, MARGIN, y, Inches(0.05), Inches(0.5), fill=color)
    tf = textbox(s, MARGIN + Inches(0.28), y + Inches(0.13), Inches(1.5), Inches(0.28))
    para(tf, ver, size=13, bold=True, color=INK, first=True, space_after=0)
    tf = textbox(s, MARGIN + Inches(1.95), y + Inches(0.14), Inches(1.0), Inches(0.26))
    para(tf, date, size=11.5, color=MUTED, first=True, space_after=0)
    tf = textbox(s, MARGIN + Inches(3.1), y + Inches(0.13), Inches(8.5), Inches(0.28))
    para(tf, note, size=12, color=INK2, first=True, space_after=0)

rect(s, MARGIN, Inches(5.42), Inches(11.83), Inches(1.05), fill=WHITE, line_color=CRITICAL,
     line_w=1.3)
rect(s, MARGIN, Inches(5.42), Inches(0.05), Inches(1.05), fill=CRITICAL)
tf = textbox(s, MARGIN + Inches(0.28), Inches(5.6), Inches(11.3), Inches(0.75))
para(tf, "对比：亚太HUB 兼职专用端  0 次发版", size=14, bold=True, color=CRITICAL,
     first=True, space_after=5)
para(tf, "版本停在 1.0.0+1、0 个 tag、无 CHANGELOG，期间还有一次「版本号重置」提交。"
         "兼职功能虽已在母 App 上线，但兼职人员专用的独立 App 尚未真正交到用户手里。",
     size=11.5, color=INK2, space_after=0, line=1.2)

# ═══ P7 独立端 ══════════════════════════════════════════════════════
s = blank(prs)
page_head(s, "亚太HUB：为兼职人员做一个专用 App", kicker="二、项目实施进度", idx=7)

tf = textbox(s, MARGIN, Inches(1.78), Inches(11.83), Inches(0.32))
para(tf, "为什么做", size=12.5, bold=True, color=BLUE, first=True, space_after=6)
tf = textbox(s, MARGIN, Inches(2.12), Inches(11.83), Inches(0.4))
para(tf, "兼职人员只需要「接单 / 打卡 / 上传素材 / 群聊」，却要装一个包含客服、商务、"
         "客户档案、日程、考勤的综合 OA —— 功能冗余、权限复杂、品牌不符。",
     size=13, color=INK2, first=True, space_after=0, line=1.25)

tf = textbox(s, MARGIN, Inches(2.72), Inches(11.83), Inches(0.32))
para(tf, "做了什么（2026-07-13 建仓，18 天内 60 个提交）", size=12.5, bold=True,
     color=BLUE, first=True, space_after=8)
items = [
    ("代码瘦身", "搬入 10.3 万行 → 砍除 5.9 万行 → 收敛为 4.5 万行；只保留兼职工作流、素材上传、任务群聊"),
    ("品牌独立", "包名 com.amgroup.hub、独立图标与启动图、隐私政策改写、移除钉钉与 Apple 登录"),
    ("推送独立", "绑定 / 登出带 app 标识，实现按 App 分别投递，兼职消息不再串到 OA"),
    ("注册门槛降低", "出生日期 / 性别 / 微信 / 韩国住址 / 护照首页 / 手机号 / 存折首页共 7 个字段改为选填"),
    ("质量守住", "净删 5.9 万行后 flutter analyze 仍为 No issues found —— 依赖、权限、平台包一并清理干净"),
]
for i, (name, desc) in enumerate(items):
    y = Inches(3.12) + i * Inches(0.62)
    rect(s, MARGIN, y, Inches(11.83), Inches(0.5), fill=WHITE, line_color=LINE)
    rect(s, MARGIN, y, Inches(0.05), Inches(0.5), fill=ORANGE)
    tf = textbox(s, MARGIN + Inches(0.26), y + Inches(0.13), Inches(2.1), Inches(0.28))
    para(tf, name, size=12.5, bold=True, color=INK, first=True, space_after=0)
    tf = textbox(s, MARGIN + Inches(2.45), y + Inches(0.13), Inches(9.2), Inches(0.28))
    para(tf, desc, size=11.5, color=INK2, first=True, space_after=0)
footnote(s, "剥离目标已达成；未收口的两点是「未正式发版」与「iOS 签名主体迁移」（见第 10 页）。")

# ═══ P8 执行成效 ════════════════════════════════════════════════════
s = blank(prs)
page_head(s, "对照立项预期：交付情况与产出价值", kicker="三、执行结果复盘", idx=8)

tf = textbox(s, MARGIN, Inches(1.78), Inches(5.7), Inches(0.32))
para(tf, "达成成效", size=13, bold=True, color=GOOD, first=True, space_after=8)
wins = [
    "6 个业务环节全部交付，闭环完整，0 个夭折",
    "从立项到闭环约 2.5 个月，是本年度从 0 到 1 最快的新业务",
    "随母 App 发版 5 次，生产环境稳定运行",
    "报销单一期内完成原定二期范围",
    "独立端剥离干净，代码健康度未受损",
]
for i, t in enumerate(wins):
    y = Inches(2.18) + i * Inches(0.62)
    rect(s, MARGIN, y, Inches(5.7), Inches(0.5), fill=WHITE, line_color=LINE)
    rect(s, MARGIN, y, Inches(0.05), Inches(0.5), fill=GOOD)
    tf = textbox(s, MARGIN + Inches(0.26), y + Inches(0.11), Inches(5.3), Inches(0.32))
    para(tf, t, size=11.5, color=INK2, first=True, space_after=0, line=1.1)

x2 = MARGIN + Inches(6.13)
tf = textbox(s, x2, Inches(1.78), Inches(5.7), Inches(0.32))
para(tf, "产出价值（已交付能力）", size=13, bold=True, color=BLUE, first=True, space_after=8)
vals = [
    ("业务链路自动化", "兼职从注册到拿钱全程在 App 内完成，"
                       "打卡、工时、交通费、报销、结算金额口径统一并与后端书面确认"),
    ("跨角色协同", "兼职与客服 / 管理员在同一会话内沟通，"
                   "并可直接查医院信息与客户资料，不需要来回问人"),
    ("专用端就绪", "兼职人员将拿到一个只装必需功能的轻量 App，"
                   "而非综合 OA"),
]
for i, (name, desc) in enumerate(vals):
    y = Inches(2.18) + i * Inches(1.05)
    rect(s, x2, y, Inches(5.7), Inches(0.92), fill=WHITE, line_color=LINE)
    rect(s, x2, y, Inches(0.05), Inches(0.92), fill=BLUE)
    tf = textbox(s, x2 + Inches(0.26), y + Inches(0.11), Inches(5.3), Inches(0.72))
    para(tf, name, size=12.5, bold=True, color=INK, first=True, space_after=4)
    para(tf, desc, size=11, color=INK2, space_after=0, line=1.15)

rect(s, MARGIN, Inches(5.62), Inches(11.83), Inches(0.82), fill=WHITE, line_color=WARNING,
     line_w=1.3)
rect(s, MARGIN, Inches(5.62), Inches(0.05), Inches(0.82), fill=WARNING)
tf = textbox(s, MARGIN + Inches(0.28), Inches(5.78), Inches(11.3), Inches(0.55))
para(tf, "必须说明：以上是「功能完成度」，不是业务收益", size=12.5, bold=True,
     color=INK, first=True, space_after=4)
para(tf, "两端均无业务埋点、独立端未发版，因此本次盘点无法提供兼职活跃人数、"
         "打卡提交成功率、结算准确率等任何业务数字。",
     size=11.5, color=INK2, space_after=0)

# ═══ P9 卡点与风险 ══════════════════════════════════════════════════
s = blank(prs)
page_head(s, "如实暴露的卡点与风险", kicker="三、执行结果复盘", idx=9)

risks = [
    ("高", CRITICAL, "独立端未正式发版",
     "停在 1.0.0+1、0 个 tag。功能与品牌均已就绪，卡在发版动作与 iOS 签名主体迁移，兼职人员仍在用综合 OA。"),
    ("高", CRITICAL, "素材上传链路零验证",
     "上传是独立端三条核心链路之一，但建仓后两端对它均零改动；新包名与新 App Group 下从未验证，含 iOS 灵动岛进度。"),
    ("高", CRITICAL, "零自动化测试",
     "两端合计 15.5 万行代码、0 个测试文件。结算金额口径、打卡工时、本地库 4 次迁移、分片上传边界全靠人工回归。"),
    ("中", WARNING, "双端漂移已实际发生",
     "母 App 于 7-31 交付「群聊图片视频改缩略图」，独立端不含该功能；共享代码靠 28% 人工同步维持，已有实际落差。"),
    ("中", WARNING, "成效无法用业务口径衡量",
     "无埋点 + 独立端未发版，导致写不出兼职活跃人数、打卡成功率、结算准确率任何一个数字。"),
    ("低", WARNING, "暂停功能无决策记录",
     "「取消任务」被整段注释屏蔽，仅留一句「暂不需要」，无暂停时间、无决策人、无恢复条件，持续占用维护成本。"),
]
for i, (level, color, title, desc) in enumerate(risks):
    y = Inches(1.8) + i * Inches(0.845)
    rect(s, MARGIN, y, Inches(11.83), Inches(0.73), fill=WHITE, line_color=LINE)
    rect(s, MARGIN, y, Inches(0.05), Inches(0.73), fill=color)
    rect(s, MARGIN + Inches(0.22), y + Inches(0.19), Inches(0.72), Inches(0.34), fill=color)
    tf = textbox(s, MARGIN + Inches(0.22), y + Inches(0.25), Inches(0.72), Inches(0.26),
                 align=PP_ALIGN.CENTER)
    para(tf, level, size=11, bold=True, color=WHITE, first=True, space_after=0,
         align=PP_ALIGN.CENTER)
    tf = textbox(s, MARGIN + Inches(1.08), y + Inches(0.08), Inches(10.5), Inches(0.56))
    para(tf, title, size=12.5, bold=True, color=INK, first=True, space_after=3)
    para(tf, desc, size=10.5, color=INK2, space_after=0, line=1.12)
footnote(s, "严重度由「对兼职业务可用性的影响 + 修复成本」评定；每条均可定位到具体文件或提交。")

# ═══ P10 下半年计划 ═════════════════════════════════════════════════
s = blank(prs)
page_head(s, "下半年推进计划（8-12 月）", kicker="四、下半年计划", idx=10)

plans = [
    ("8 月", "让独立端真正上线",
     "① 补齐素材上传六场景验证（直传 / 分片 / 并发 / 断网重试 / 前后台恢复 / 灵动岛）\n"
     "② 打首个 tag v1.0.0 + 建 CHANGELOG   ③ 推进 iOS 签名主体迁移", CRITICAL, "最高"),
    ("8 月", "止住双端漂移",
     "① 立即把母 App 的群聊缩略图功能同步到独立端\n"
     "② 建「双端共享文件清单」+ 差异校验脚本 —— 不求自动同步，只求漂移能被发现", CRITICAL, "紧急"),
    ("8-9 月", "偿还技术债",
     "① 对「取消任务」作出恢复或删除的明确决策   ② 扫描并清理全部过期「二期待接」标记\n"
     "③ 建立后端阻塞台账，登记前端就绪、卡契约的条目", WARNING, "高"),
    ("9 月起", "质量与数据前置",
     "① 为结算金额、打卡工时、本地库迁移三处关键逻辑补单元测试，接入 CI 门禁\n"
     "② 上最小埋点四个数：打卡提交成功率、上传成功率、消息送达率、注册通过率", WARNING, "常态"),
]
for i, (when, title, desc, color, pri) in enumerate(plans):
    y = Inches(1.82) + i * Inches(1.22)
    rect(s, MARGIN, y, Inches(11.83), Inches(1.08), fill=WHITE, line_color=LINE)
    rect(s, MARGIN, y, Inches(0.05), Inches(1.08), fill=color)
    tf = textbox(s, MARGIN + Inches(0.26), y + Inches(0.12), Inches(1.1), Inches(0.28))
    para(tf, when, size=13, bold=True, color=color, first=True, space_after=0)
    tf = textbox(s, MARGIN + Inches(1.5), y + Inches(0.11), Inches(7.0), Inches(0.3))
    para(tf, title, size=13.5, bold=True, color=INK, first=True, space_after=0)
    rect(s, Inches(11.6), y + Inches(0.13), Inches(0.92), Inches(0.3), fill=color)
    tf = textbox(s, Inches(11.6), y + Inches(0.18), Inches(0.92), Inches(0.26),
                 align=PP_ALIGN.CENTER)
    para(tf, pri, size=10.5, bold=True, color=WHITE, first=True, space_after=0,
         align=PP_ALIGN.CENTER)
    tf = textbox(s, MARGIN + Inches(1.5), y + Inches(0.46), Inches(9.9), Inches(0.56))
    for j, ln in enumerate(desc.split("\n")):
        para(tf, ln, size=11, color=INK2, first=(j == 0), space_after=2, line=1.1)

rect(s, MARGIN, Inches(6.72), Inches(11.83), Inches(0.4), fill=WHITE, line_color=LINE)
tf = textbox(s, MARGIN + Inches(0.26), Inches(6.8), Inches(11.3), Inches(0.26))
para(tf, "第五优先级：新项目储备 —— 两个仓库均无下半年兼职新项目立项信息，需业务侧 8 月中旬前明确。",
     size=11, bold=True, color=INK2, first=True, space_after=0)

# ═══ P11 需要决策与支持 ═════════════════════════════════════════════
s = blank(prs)
page_head(s, "需要管理层决策与支持的事项", kicker="四、下半年计划", idx=11)

asks = [
    ("1", "独立端发版时间与 iOS 签名主体迁移",
     "亚太HUB 功能已就绪，发版需确定时间窗口；iOS 签名主体迁移属外部后台事项，需管理层推动。",
     "需明确：发版目标日期 + 签名迁移责任人"),
    ("2", "提供兼职业务的量化指标",
     "本次盘点无立项书、无 KPI、无排期基线，「是否达标」只能以反推口径近似判断。",
     "需提供：兼职人数目标、打卡 / 结算准确率要求等可验证标准"),
    ("3", "「取消任务」功能的去留",
     "该功能被屏蔽且无决策记录，代码完整保留在两个仓库中持续占用维护成本。",
     "需决策：恢复上线 or 明确删除"),
    ("4", "研发单点依赖",
     "本年度 93.2% 的提交来自同一名开发者，兼职业务的核心知识高度集中，无人员冗余。",
     "需评估：是否补充人力 or 通过评审与文档分散知识"),
    ("5", "保障性投入的排期认可",
     "补测试、补埋点、上传链路验证均不产生新功能，但决定了兼职业务能否被衡量与安全迭代。",
     "需认可：为这些工作预留正式排期，而非用零碎时间挤"),
]
for i, (num, title, why, ask) in enumerate(asks):
    y = Inches(1.8) + i * Inches(1.02)
    rect(s, MARGIN, y, Inches(11.83), Inches(0.88), fill=WHITE, line_color=LINE)
    rect(s, MARGIN, y, Inches(0.05), Inches(0.88), fill=BLUE)
    rect(s, MARGIN + Inches(0.24), y + Inches(0.26), Inches(0.36), Inches(0.36),
         fill=BLUE, shape=MSO_SHAPE.OVAL)
    tf = textbox(s, MARGIN + Inches(0.24), y + Inches(0.32), Inches(0.36), Inches(0.26),
                 align=PP_ALIGN.CENTER)
    para(tf, num, size=11.5, bold=True, color=WHITE, first=True, space_after=0,
         align=PP_ALIGN.CENTER)
    tf = textbox(s, MARGIN + Inches(0.78), y + Inches(0.1), Inches(6.4), Inches(0.7))
    para(tf, title, size=13, bold=True, color=INK, first=True, space_after=4)
    para(tf, why, size=10.5, color=INK2, space_after=0, line=1.12)
    rect(s, Inches(8.5), y + Inches(0.14), Inches(4.0), Inches(0.6), fill=SURFACE,
         line_color=LINE)
    tf = textbox(s, Inches(8.66), y + Inches(0.24), Inches(3.7), Inches(0.42))
    para(tf, ask, size=10.5, bold=True, color=BLUE, first=True, space_after=0, line=1.12)

# ═══ P12 数据口径 ═══════════════════════════════════════════════════
s = blank(prs)
page_head(s, "附：数据口径与边界说明", kicker="附页", idx=12)

tf = textbox(s, MARGIN, Inches(1.8), Inches(11.83), Inches(0.3))
para(tf, "数据来源", size=12.5, bold=True, color=BLUE, first=True, space_after=6)
tf = textbox(s, MARGIN, Inches(2.12), Inches(11.83), Inches(0.4))
para(tf, "flutter-suixinpai（2026 年 663 个提交）与 yatai_hub（60 个提交）两个仓库的 git 记录、"
         "pubspec.yaml 版本记录、分支状态、源码现状与静态分析结果。",
     size=12, color=INK2, first=True, space_after=0, line=1.2)

tf = textbox(s, MARGIN, Inches(2.72), Inches(11.83), Inches(0.3))
para(tf, "三点边界（避免误读）", size=12.5, bold=True, color=CRITICAL, first=True,
     space_after=8)
bounds = [
    ("「预设指标」不是立项时的书面 KPI",
     "两个仓库均无立项书、无 OKR、无排期基线。本报告的交付口径系依据首个提交范围与发版节点反推，"
     "用于衡量「是否交付完整」，不代表管理层签署的量化目标。"),
    ("「产出价值」只描述已交付能力，不含业务收益数字",
     "「提效 X%」「节省 Y 工时」一类数据代码仓库无从得知，本报告一概不写。"),
    ("「项目」为归纳单位",
     "仓库以功能提交为最小粒度，无显式项目编号。P08 / P15 编号系本次盘点新拟，便于后续台账跟踪。"),
]
for i, (t, d) in enumerate(bounds):
    y = Inches(3.12) + i * Inches(0.92)
    rect(s, MARGIN, y, Inches(11.83), Inches(0.78), fill=WHITE, line_color=LINE)
    rect(s, MARGIN, y, Inches(0.05), Inches(0.78), fill=MUTED)
    tf = textbox(s, MARGIN + Inches(0.26), y + Inches(0.11), Inches(11.3), Inches(0.6))
    para(tf, t, size=12, bold=True, color=INK, first=True, space_after=4)
    para(tf, d, size=10.5, color=INK2, space_after=0, line=1.15)

rect(s, MARGIN, Inches(5.92), Inches(11.83), Inches(0.82), fill=SURFACE, line_color=BLUE)
tf = textbox(s, MARGIN + Inches(0.26), Inches(6.06), Inches(11.3), Inches(0.6))
para(tf, "完整版盘点报告", size=12, bold=True, color=BLUE, first=True, space_after=4)
para(tf, "覆盖全部 15 个立项项目（含客服、商务 BD、考勤、审批等非兼职模块）的双端整合版报告，"
         "见 yatai_hub/docs/2026年度项目盘点报告（双端整合版）.md",
     size=10.5, color=INK2, space_after=0)

out = "/Users/tianjian/Desktop/project/yatai_hub/docs/2026兼职业务盘点汇报.pptx"
prs.save(out)
print(f"已生成: {out}")
print(f"页数: {len(prs.slides.__iter__.__self__._sldIdLst)}")
